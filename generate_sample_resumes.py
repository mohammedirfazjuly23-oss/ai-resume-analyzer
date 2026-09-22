"""
generate_sample_resumes.py - Utility to generate sample PDF, DOCX, and TXT resumes
"""

import os
import docx
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

SAMPLE_DIR = os.path.join(os.path.dirname(__file__), 'sample_resumes')
os.makedirs(SAMPLE_DIR, exist_ok=True)

# 1. Data Analyst Sample TXT Resume
DATA_ANALYST_TXT = """Alex Johnson
Data Analyst | Business Intelligence Specialist
Email: alex.johnson@example.com | Phone: (555) 123-4567

OBJECTIVE
Motivated Data Analyst with 2+ years of experience transforming complex datasets into actionable business insights.

SKILLS
- Programming: Python, SQL, R
- Data Analysis & Tools: Pandas, NumPy, Excel, Power BI, Tableau
- Databases: PostgreSQL, MySQL
- Methodologies: Statistics, Data Visualization, ETL pipelines

WORK EXPERIENCE
Data Analyst Intern | Insights Corp (2023 - Present)
- Extracted and cleaned data using Python and Pandas for 15+ marketing campaigns.
- Designed interactive Power BI dashboards to track KPIs, improving reporting efficiency by 30%.
- Queried PostgreSQL databases to write complex SQL joins and aggregations.

PROJECTS
Sales Performance Analytics Dashboard
- Built a Power BI interactive report consuming sales data from Excel and SQL.
- Applied statistical forecasting to predict quarterly sales with 88% accuracy.

EDUCATION
B.S. in Statistics and Computer Science | State University (2023)
"""

# 2. ML Engineer DOCX Resume
def create_ml_docx():
    file_path = os.path.join(SAMPLE_DIR, 'sample_resume_ml_engineer.docx')
    doc = docx.Document()
    doc.add_heading('Samira Patel', level=0)
    doc.add_paragraph('Machine Learning Engineer | AI Developer\nEmail: samira.patel@example.com')
    
    doc.add_heading('SUMMARY', level=1)
    doc.add_paragraph('Enthusiastic ML Engineer with background in building predictive models, scikit-learn pipelines, and Python backend services.')

    doc.add_heading('TECHNICAL SKILLS', level=1)
    doc.add_paragraph('Languages: Python, C++, SQL\nML & AI: Machine Learning, scikit-learn, Deep Learning, PyTorch, Pandas, NumPy\nTools & Cloud: Git, Linux, Docker')

    doc.add_heading('EXPERIENCE', level=1)
    p = doc.add_paragraph()
    p.add_run('Junior ML Engineer - AI Solutions Inc.\n').bold = True
    p.add_run('• Trained predictive machine learning models using scikit-learn and XGBoost.\n')
    p.add_run('• Processed structured tabular datasets using Python, Pandas, and SQL.\n')
    p.add_run('• Containerized ML scripts using Docker for deployment.')

    doc.add_heading('PROJECTS', level=1)
    doc.add_paragraph('Customer Churn Classifier: Developed a classification pipeline using scikit-learn with 91% ROC-AUC score.')

    doc.add_heading('EDUCATION', level=1)
    doc.add_paragraph('B.Tech in Computer Science & Engineering (2023)')
    doc.save(file_path)
    return file_path


# 3. NLP Engineer PDF Resume
def create_nlp_pdf():
    file_path = os.path.join(SAMPLE_DIR, 'sample_resume_nlp_engineer.pdf')
    doc = SimpleDocTemplate(file_path, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    
    story = []
    story.append(Paragraph("David Chen", styles['Heading1']))
    story.append(Paragraph("NLP Engineer | AI Researcher | david.chen@example.com", styles['Normal']))
    story.append(Spacer(1, 10))
    
    story.append(Paragraph("SKILLS", styles['Heading2']))
    story.append(Paragraph("Python, NLP, Transformers, Hugging Face, spaCy, NLTK, PyTorch, BERT, Deep Learning, REST API, Git", styles['Normal']))
    story.append(Spacer(1, 10))

    story.append(Paragraph("EXPERIENCE", styles['Heading2']))
    story.append(Paragraph("<b>NLP Engineer Intern</b> - Text AI Labs (2023 - Present)", styles['Normal']))
    story.append(Paragraph("• Fine-tuned Hugging Face Transformers (BERT, RoBERTa) for sentiment analysis and text classification.", styles['Normal']))
    story.append(Paragraph("• Built custom Named Entity Recognition (NER) pipelines using spaCy and Python.", styles['Normal']))
    story.append(Paragraph("• Processed large text datasets with NLTK and tokenization libraries.", styles['Normal']))
    story.append(Spacer(1, 10))

    story.append(Paragraph("EDUCATION", styles['Heading2']))
    story.append(Paragraph("M.S. in Computational Linguistics | Tech University", styles['Normal']))
    
    doc.build(story)
    return file_path


if __name__ == '__main__':
    # Write TXT
    txt_path = os.path.join(SAMPLE_DIR, 'sample_resume_data_analyst.txt')
    with open(txt_path, 'w', encoding='utf-8') as f:
        f.write(DATA_ANALYST_TXT)
    
    docx_path = create_ml_docx()
    pdf_path = create_nlp_pdf()

    print("Sample resumes generated successfully:")
    print(" -", txt_path)
    print(" -", docx_path)
    print(" -", pdf_path)
